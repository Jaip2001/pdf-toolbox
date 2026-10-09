"""
converter.py — PDF (jewellery estimate / any ruled table) -> Excel

Two engines:
  * estimate : the exact "ESTIMATE ... SJ3910" style bill layout (same-to-same replica)
  * generic  : any PDF that has ruled tables (structure preserved)

Also extracts every image of the PDF into a zip.
"""
import os
import re
import io
import json
import math
import shutil
import zipfile
import hashlib
from collections import Counter, defaultdict

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.pagebreak import Break

EMU_PER_PT = 12700
THIN = Side(style='thin', color='FF000000')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FN = 'Calibri'
F = Font(name=FN, size=7)
FB = Font(name=FN, size=7, bold=True)
FI = Font(name=FN, size=7, italic=True)
F8B = Font(name=FN, size=8, bold=True)
R = Alignment(horizontal='right', vertical='center', shrink_to_fit=True)
LT = Alignment(horizontal='left', vertical='top', wrap_text=True)
LC = Alignment(horizontal='left', vertical='center')
C = Alignment(horizontal='center', vertical='center', wrap_text=True)

INR = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
FMT_WT = '0.000'
FMT_PCS = '0'
FMT_MONEY = '#,##0.00'

# =====================================================================
#  small helpers
# =====================================================================


def clean(t):
    return re.sub(r'[ \t]+', ' ', t.replace('\xa0', ' ')).strip()


def is_num(s):
    return bool(re.fullmatch(r'-?[\d,]+\.?\d*', s.strip())) and s.strip() not in ('', '.', '-')


def num(s):
    try:
        return float(s.replace(',', ''))
    except Exception:
        return s


def fix_wrapped_number(txt):
    """'1,06,026.0\\n0'  ->  '1,06,026.00'"""
    parts = [p.strip() for p in txt.split('\n') if p.strip()]
    if len(parts) > 1 and all(re.fullmatch(r'\d[\d,]*\.?\d*', p) for p in parts):
        joined = ''.join(parts)
        if is_num(joined):
            return joined
    return txt


def number_format_for(text):
    """match the look of the pdf cell (decimals / grouping)"""
    t = text.strip()
    if not is_num(t):
        return None
    dec = len(t.split('.')[1]) if '.' in t else 0
    if dec == 0:
        return FMT_PCS
    if dec == 3:
        return FMT_WT
    return FMT_MONEY


def excel_col_width(pt):
    return max(round(((pt * 96 / 72) - 5) / 7, 2), 3.5)


def put_cell(ws, r, c, value, font=F, align=R, fmt=None):
    cell = ws.cell(row=r, column=c)
    if isinstance(value, str):
        v = value.strip()
        if v == '':
            cell.value = None
            cell.font = font
            if align:
                cell.alignment = align
            return cell
        if is_num(v):
            cell.value = num(v)
            cell.number_format = fmt or number_format_for(v) or 'General'
        else:
            cell.value = value
    else:
        cell.value = value
        if fmt:
            cell.number_format = fmt
    cell.font = font
    if align:
        cell.alignment = align
    return cell


# =====================================================================
#  template detection
# =====================================================================


def detect_template(doc):
    """returns 'estimate' | 'generic' | 'text'"""
    try:
        page = doc[0]
    except Exception:
        return 'text'
    txt = page.get_text()
    if 'ESTIMATE' not in txt.upper()[:400]:
        return table_check(doc)
    tabs = page.find_tables()
    if not tabs.tables:
        return 'text'
    t = tabs.tables[0]
    try:
        sub = page.get_textbox(pymupdf.Rect(t.rows[1].bbox)).replace('\n', ' ')
    except Exception:
        return table_check(doc)
    need = ['Code', 'Size', 'Pcs', 'Wt', 'Rate', 'Amount', 'Quality', 'Net Wt']
    if all(k in sub for k in need):
        edges = edges_of(t)
        if len(edges) == 24:            # 23 columns, exactly like the SJ3910 bill
            return 'estimate'
    return 'generic'


def table_check(doc):
    for pno in range(min(len(doc), 3)):
        if doc[pno].find_tables().tables:
            return 'generic'
    return 'text'


def edges_of(table):
    """all distinct x edges (lefts + rights) of the first two rows of a table"""
    xs = set()
    for ri in (0, 1):
        if ri >= table.row_count:
            continue
        for c in table.rows[ri].cells:
            if c:
                xs.add(round(c[0], 1))
                xs.add(round(c[2], 1))
    return sorted(xs)


# =====================================================================
#  image helpers
# =====================================================================


def collect_page_images(doc, pno, out_dir, prefix, registry=None):
    """save the page images, return list of dicts"""
    page = doc[pno]
    out = []
    for i, info in enumerate(page.get_image_info(xrefs=True)):
        xref = info['xref']
        if not xref:
            continue
        ext = doc.extract_image(xref)
        if not ext or not ext.get('image'):
            continue
        data = ext['image']
        h = hashlib.md5(data).hexdigest()[:12]
        name = f'{h}.{ext["ext"]}'
        path = os.path.join(out_dir, name)
        if not os.path.exists(path):
            with open(path, 'wb') as fh:
                fh.write(data)
        out.append({'path': path, 'name': name, 'bbox': list(info['bbox']),
                    'w': ext['width'], 'h': ext['height'], 'page': pno,
                    'ext': ext['ext'], 'hash': h})
    return out


# =====================================================================
#  ESTIMATE engine  (exact replica)
# =====================================================================

COLS = ['sr', 'design', 'd_code', 'd_size', 'd_pcs', 'd_wt', 'd_rate', 'd_amt',
        'm_quality', 'm_wt', 'm_netwt', 'm_rate', 'm_amt',
        's_code', 's_size', 's_pcs', 's_wt', 's_rate', 's_amt',
        'o_amt', 'l_rate', 'l_amt', 'total']
NUMCOLS = ['d_pcs', 'd_wt', 'd_rate', 'd_amt', 'm_wt', 'm_netwt', 'm_rate', 'm_amt',
           's_pcs', 's_wt', 's_rate', 's_amt', 'o_amt', 'l_rate', 'l_amt', 'total']


def parse_header(doc):
    """To-name, invoice number and date of an estimate bill"""
    out = {'to': '', 'invoice': '', 'date': '', 'title': 'ESTIMATE'}
    try:
        spans = []
        for b in doc[0].get_text('dict')['blocks']:
            if b['type'] != 0:
                continue
            for l in b['lines']:
                for s in l['spans']:
                    t = clean(s['text'])
                    if t:
                        spans.append({'t': t, 'bbox': list(s['bbox'])})
        for s in spans:
            x0, y0, x1, y1 = s['bbox']
            if s['t'].startswith('Invoice'):
                nxt = [q for q in spans if q['bbox'][0] > x1 - 1 and abs(q['bbox'][1] - y0) < 4]
                if nxt:
                    out['invoice'] = nxt[0]['t']
            elif s['t'].startswith('Date'):
                nxt = [q for q in spans if q['bbox'][0] > x1 - 1 and abs(q['bbox'][1] - y0) < 4]
                if nxt:
                    out['date'] = nxt[0]['t']
            elif s['t'] == 'To':
                below = [q for q in spans if q['bbox'][1] >= y1 - 4 and q['bbox'][1] < y1 + 16
                         and q['bbox'][0] < 320 and q['t'] != 'To']
                if below:
                    below.sort(key=lambda q: q['bbox'][1])
                    out['to'] = below[0]['t']
        # title line
        top = [s for s in spans if s['bbox'][1] < 40]
        if top:
            out['title'] = top[0]['t'].strip()
    except Exception:
        pass
    return out


def parse_estimate(doc, img_dir):
    pages = []
    for pno in range(len(doc)):
        page = doc[pno]
        spans, imgs, bands = [], [], []
        for b in page.get_text('dict')['blocks']:
            if b['type'] == 1:
                imgs.append({'bbox': list(b['bbox']), 'pno': pno, 'size': b['size']})
                continue
            for l in b['lines']:
                for s in l['spans']:
                    t = clean(s['text'])
                    if t:
                        spans.append({'t': t, 'bbox': list(s['bbox'])})
        tabs = page.find_tables()
        if tabs.tables:
            bands = [list(r.bbox) for r in tabs.tables[0].rows]
        pages.append({'spans': spans, 'imgs': imgs, 'bands': bands, 'no': pno})

    # dynamic column boundaries from page 1 header
    t0 = doc[0].find_tables().tables[0]
    X = edges_of(t0)
    if len(X) != 24:
        raise ValueError('estimate layout could not be read')

    def col_of(x0, x1):
        best, bi = -1, 1
        for i in range(len(X) - 1):
            ov = min(x1, X[i + 1]) - max(x0, X[i])
            if ov > best:
                best, bi = ov, i
        return bi

    # ---- item anchors
    starts = []
    for pg in pages:
        for s in pg['spans']:
            x0, y0, x1, y1 = s['bbox']
            if x0 < X[1] + 0.3:
                m = re.match(r'^(\d+)\s*(.*)$', s['t'])
                if m and (x0 < X[1] - 7 or x1 <= X[1] + 1.4):
                    starts.append({'pg': pg['no'], 'y': round(y0, 1),
                                   'sr': int(m.group(1)), 'rest': m.group(2).strip()})
    starts.sort(key=lambda a: (a['pg'], a['y']))
    ded = []
    for st in starts:
        if ded and ded[-1]['pg'] == st['pg'] and abs(ded[-1]['y'] - st['y']) < 2:
            continue
        ded.append(st)
    starts = ded

    last_pg = pages[-1]
    footer_cut = None
    for s in last_pg['spans']:
        if s['t'].strip() == 'Total' and s['bbox'][1] > 250:
            footer_cut = s['bbox'][1] - 5
    for s in last_pg['spans']:
        if s['t'].strip() == 'Add' and s['bbox'][0] > 700:
            footer_cut = min(footer_cut if footer_cut else 1e9, s['bbox'][1] - 5)

    items = []
    for i, st in enumerate(starts):
        pg = pages[st['pg']]
        y_end = 1e9
        if i + 1 < len(starts) and starts[i + 1]['pg'] == st['pg']:
            y_end = starts[i + 1]['y']
        if st['pg'] == len(pages) - 1 and footer_cut:
            y_end = min(y_end, footer_cut)
        y0 = st['y'] - 1.5

        spans = [s for s in pg['spans'] if y0 <= s['bbox'][1] < y_end - 1.5]
        cells = {c: [] for c in COLS}
        for s in spans:
            x0_, y0_, x1_, y1_ = s['bbox']
            c = col_of(x0_, x1_)
            if c == 1 and x0_ < X[1] + 0.3:
                c = 0
            cells[COLS[c]].append({'t': s['t'], 'x': round(x0_, 1), 'y': round(y0_, 1),
                                   'bbox': [x0_, y0_, x1_, y1_]})

        # design block
        lines, cur = [], []
        for s in sorted(cells['design'], key=lambda d: (d['y'], d['x'])):
            if cur and abs(s['y'] - cur[-1]['y']) > 2:
                lines.append(cur)
                cur = []
            cur.append(s)
        if cur:
            lines.append(cur)
        design_name = code = colour = tunch = gross = ''
        for ln in lines:
            txt = ' '.join(d['t'] for d in ln)
            if not txt:
                continue
            y = ln[0]['y']
            if txt.startswith('Tunch'):
                tunch = txt
            elif 'gm' in txt and 'Gross' in txt:
                gross = txt
            elif y <= st['y'] + 1.2:
                parts = [t['t'] for t in ln if t['x'] < 60]
                cds = [t['t'] for t in ln if t['x'] >= 60]
                if parts:
                    design_name = ' '.join(parts).strip()
                if cds:
                    code = ' '.join(cds).strip()
            elif not colour:
                colour = txt

        # subtotal cut
        total_ys = [s['bbox'][1] for s in spans
                    if col_of(s['bbox'][0], s['bbox'][2]) == COLS.index('total')
                    and is_num(s['t'])]
        sub_top = max(total_ys) - 9 if total_ys else y_end - 22

        d_anchor = sorted({s['y'] for s in cells['d_rate'] if is_num(s['t'])})
        s_anchor = sorted({s['y'] for s in cells['s_rate'] if is_num(s['t'])})

        def assign(col, anchors):
            out = [[] for _ in anchors]
            rest = []
            for s in sorted(cells[col], key=lambda d: d['y']):
                if s['y'] >= sub_top:
                    rest.append(s)
                    continue
                k = 0
                for j in range(len(anchors)):
                    if s['y'] >= anchors[j] - 0.6:
                        k = j
                out[k].append(s)
            return out, rest

        def val(lines_):
            if not lines_:
                return ''
            lines_ = sorted(lines_, key=lambda d: d['y'])
            if len(lines_) == 1:
                return num(lines_[0]['t'])
            j = ''.join(d['t'] for d in lines_)
            if is_num(j):
                return num(j)
            return '\n'.join(d['t'] for d in lines_)

        diamonds, stones, sub_pool = [], [], {'d': [], 's': []}
        for col in ('d_code', 'd_size', 'd_pcs', 'd_wt', 'd_rate', 'd_amt'):
            per, rest = assign(col, d_anchor)
            for j, ls in enumerate(per):
                while len(diamonds) <= j:
                    diamonds.append({})
                diamonds[j][col] = val(ls)
            sub_pool['d'] += rest
        for col in ('s_code', 's_size', 's_pcs', 's_wt', 's_rate', 's_amt'):
            per, rest = assign(col, s_anchor)
            for j, ls in enumerate(per):
                while len(stones) <= j:
                    stones.append({})
                stones[j][col] = val(ls)
            sub_pool['s'] += rest

        item_level, sub_rest = {}, {}
        for col in ('m_quality', 'm_wt', 'm_netwt', 'm_rate', 'm_amt',
                    'o_amt', 'l_rate', 'l_amt', 'total'):
            main = [s for s in cells[col] if s['y'] < sub_top]
            rest = [s for s in cells[col] if s['y'] >= sub_top]
            item_level[col] = val(main)
            sub_rest[col] = val(rest)
        for col in ('d_pcs', 'd_wt', 'd_rate', 'd_amt'):
            sub_rest[col] = val([s for s in sub_pool['d']
                                 if col_of(s['bbox'][0], s['bbox'][2]) == COLS.index(col)])
        for col in ('s_pcs', 's_wt', 's_rate', 's_amt'):
            sub_rest[col] = val([s for s in sub_pool['s']
                                 if col_of(s['bbox'][0], s['bbox'][2]) == COLS.index(col)])

        imgs = [im for im in pg['imgs']
                if X[1] - 1 <= im['bbox'][0] <= X[2] and y0 <= im['bbox'][1] < y_end - 1.5]

        items.append({
            'sr': st['sr'], 'page': st['pg'] + 1,
            'design': design_name or st.get('rest', ''), 'code': code, 'colour': colour,
            'tunch': tunch, 'gross': gross,
            'diamonds': diamonds, 'stones': stones, 'item': item_level, 'sub': sub_rest,
            'image': imgs[0] if imgs else None,
        })

    # map images
    page_cache = {}
    for it in items:
        if not it['image']:
            continue
        pno = it['image']['pno']
        if pno not in page_cache:
            page_cache[pno] = collect_page_images(doc, pno, img_dir, '')
        best, bd = None, 1e9
        for im in page_cache[pno]:
            d = abs(im['bbox'][0] - it['image']['bbox'][0]) + abs(im['bbox'][1] - it['image']['bbox'][1])
            if d < bd:
                bd, best = d, im
        if best and bd < 3:
            it['image_file'] = best['path']
    return {'items': items, 'X': X}


def build_estimate_xlsx(parsed, out_path, with_images=True):
    items = parsed['items']
    X = parsed['X']
    hdr = parsed.get('header') or {}
    pt_widths = [X[i + 1] - X[i] for i in range(len(X) - 1)]

    wb = Workbook()
    ws = wb.active
    ws.title = 'ESTIMATE'
    widths = {1: 3.5}
    widths[2] = min(max(excel_col_width(pt_widths[1]), 12.0), 30.0)   # Design column
    for i, pt in zip(range(3, 24), pt_widths[2:]):
        widths[i] = excel_col_width(pt)
    # keep numbers readable
    widths[7] = max(widths[7], 8.8)
    widths[8] = max(widths[8], 10.4)
    widths[18] = max(widths[18], 8.0)
    widths[19] = max(widths[19], 8.6)
    widths[21] = max(widths[21], 7.4)
    widths[22] = max(widths[22], 9.4)
    widths[23] = max(widths[23], 10.8)
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    NC = 23
    widths[20] = max(widths[20], 11.5)
    widths[9] = max(widths[9], 8.6)      # Quality (GOLD 14K / 80)
    widths[14] = max(widths[14], 12.6)   # Stone code (COLOUR STONE 100)
    COLFMT = {5: FMT_PCS, 6: FMT_WT, 7: INR, 8: INR,
              10: FMT_WT, 11: FMT_WT, 12: INR, 13: INR,
              16: FMT_PCS, 17: FMT_WT, 18: INR, 19: INR,
              20: INR, 21: INR, 22: INR, 23: INR}
    DESIGN_W_PT = widths[2] * 7 * 72 / 96

    def colw_pt(c):
        return widths[c] * 7 * 72 / 96

    # ---- banner
    head = items[0]['page'] if items else 1
    top_txt = None
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NC)
    ws.cell(row=1, column=1, value=(hdr.get('title') or 'ESTIMATE') + ' ').font = F8B
    ws.cell(row=1, column=1).alignment = LC
    for c in range(1, NC + 1):
        ws.cell(row=1, column=c).border = Border(left=THIN if c == 1 else None,
                                                 right=THIN if c == NC else None,
                                                 top=THIN, bottom=THIN)
    ws.row_dimensions[1].height = 12
    ws.cell(row=2, column=1, value='To').font = F
    ws.cell(row=3, column=1, value=hdr.get('to', '')).font = Font(name=FN, size=8, bold=True)
    ws.cell(row=2, column=21, value='Invoice# :').font = FB
    ws.cell(row=2, column=21).alignment = R
    ws.merge_cells(start_row=2, start_column=22, end_row=2, end_column=23)
    ws.cell(row=2, column=22, value=hdr.get('invoice', '')).font = FB
    ws.cell(row=2, column=22).alignment = LC
    ws.cell(row=3, column=21, value='Date :').font = FB
    ws.cell(row=3, column=21).alignment = R
    ws.merge_cells(start_row=3, start_column=22, end_row=3, end_column=23)
    ws.cell(row=3, column=22, value=hdr.get('date', '')).font = FB
    ws.cell(row=3, column=22).alignment = LC
    ws.row_dimensions[2].height = 11
    ws.row_dimensions[3].height = 11
    ws.row_dimensions[4].height = 3
    for r in (2, 3, 4):
        ws.row_dimensions[r].height = ws.row_dimensions[r].height or 11

    # ---- table head
    HDR1, HDR2 = 5, 6
    for c1, c2, label in [(3, 8, 'Diamond'), (9, 13, 'Metal'), (14, 19, 'Stone & Misc'),
                          (21, 22, 'Labour')]:
        ws.merge_cells(start_row=HDR1, start_column=c1, end_row=HDR1, end_column=c2)
        ws.cell(row=HDR1, column=c1, value=label).font = FB
        ws.cell(row=HDR1, column=c1).alignment = C
    for c1, c2, label in [(1, 1, 'Sr'), (2, 2, 'Design'), (20, 20, 'Other')]:
        ws.merge_cells(start_row=HDR1, start_column=c1, end_row=HDR2, end_column=c2)
        ws.cell(row=HDR1, column=c1, value=label).font = FB
        ws.cell(row=HDR1, column=c1).alignment = C
    ws.cell(row=HDR1, column=23, value='Total\nAmount').font = FB
    ws.cell(row=HDR1, column=23).alignment = C
    for col, label in {3: 'Code', 4: 'Size', 5: 'Pcs', 6: 'Wt', 7: 'Rate', 8: 'Amount',
                       9: 'Quality', 10: '*Wt', 11: 'Net Wt', 12: 'Rate', 13: 'Amount',
                       14: 'Code', 15: 'Size', 16: 'Pcs', 17: 'Wt', 18: 'Rate',
                       19: 'Amount', 21: 'Rate', 22: 'Amount'}.items():
        ws.cell(row=HDR2, column=col, value=label).font = FB
        ws.cell(row=HDR2, column=col).alignment = C
    for r in (HDR1, HDR2):
        for c in range(1, NC + 1):
            ws.cell(row=r, column=c).border = BORDER
    ws.row_dimensions[HDR1].height = 19.5   # 'Total Amount' is a 2-line header
    ws.row_dimensions[HDR2].height = 12

    # ---- body
    ROWH_A, ROWH_B, ROWH_C, ROW_H_SUB, IMG_PT = 30.5, 41.5, 19.0, 16.0, 41.5
    row = 7
    page_rows = {}
    for it in items:
        nd = max(len(it['diamonds']), 1)
        ns = len(it['stones'])
        first = row
        rA, rB, rC = first, first + 1, first + 2
        last = first + 3
        ws.row_dimensions[rA].height = ROWH_A
        ws.row_dimensions[rB].height = ROWH_B
        ws.row_dimensions[rC].height = ROWH_C
        ws.row_dimensions[last].height = ROW_H_SUB
        line_rows = [rA, rB, rC]

        ws.merge_cells(start_row=first, start_column=1, end_row=last, end_column=1)
        ws.cell(row=first, column=1, value=it['sr']).font = F
        ws.cell(row=first, column=1).alignment = Alignment(horizontal='left', vertical='top')
        c = ws.cell(row=rA, column=2,
                    value=f"{it['design']}\n{it['code']}\n{it['colour']}")
        c.font = F
        c.alignment = LT
        c = ws.cell(row=rC, column=2, value=f"{it['tunch']}\n{it['gross']}")
        c.font = F
        c.alignment = LT

        for i in range(nd):
            d = it['diamonds'][i] if i < len(it['diamonds']) else {}
            r = line_rows[i]
            for col, key in [(3, 'd_code'), (4, 'd_size'), (5, 'd_pcs'),
                             (6, 'd_wt'), (7, 'd_rate'), (8, 'd_amt')]:
                v = d.get(key, '')
                if isinstance(v, str) and v == '':
                    ws.cell(row=r, column=col).font = F
                    continue
                put_cell(ws, r, col, v, F, LT if key in ('d_code', 'd_size') else R,
                         fmt=COLFMT.get(col))
        for i in range(ns):
            s = it['stones'][i]
            r = line_rows[i]
            for col, key in [(14, 's_code'), (15, 's_size'), (16, 's_pcs'),
                             (17, 's_wt'), (18, 's_rate'), (19, 's_amt')]:
                v = s.get(key, '')
                if isinstance(v, str) and v == '':
                    ws.cell(row=r, column=col).font = F
                    continue
                put_cell(ws, r, col, v, F, LT if key in ('s_code', 's_size') else R,
                         fmt=COLFMT.get(col))

        for col in (9, 10, 11, 12, 13, 20, 21, 22, 23):
            ws.merge_cells(start_row=rA, start_column=col, end_row=rC, end_column=col)
        m = it['item']
        for col, key in [(9, 'm_quality'), (10, 'm_wt'), (11, 'm_netwt'), (12, 'm_rate'),
                         (13, 'm_amt'), (20, 'o_amt'), (21, 'l_rate'), (22, 'l_amt'),
                         (23, 'total')]:
            v = m.get(key, '')
            if isinstance(v, str) and v == '':
                ws.cell(row=rA, column=col).font = F
                continue
            put_cell(ws, rA, col, v, F, LT if key in ('o_amt', 'm_quality') else R,
                     fmt=COLFMT.get(col))

        s2 = it['sub']
        for col, key in [(5, 'd_pcs'), (6, 'd_wt'), (8, 'd_amt'), (10, 'm_wt'),
                         (11, 'm_netwt'), (13, 'm_amt'), (16, 's_pcs'), (17, 's_wt'),
                         (18, 's_rate'), (19, 's_amt'), (20, 'o_amt'), (22, 'l_amt'),
                         (23, 'total')]:
            v = s2.get(key, '')
            if isinstance(v, str) and v == '':
                ws.cell(row=last, column=col).font = F
                continue
            put_cell(ws, last, col, v, F, R, fmt=COLFMT.get(col))

        if with_images and it.get('image_file') and os.path.exists(it['image_file']):
            img = XLImage(it['image_file'])
            size_px = IMG_PT * 96 / 72
            img.width = size_px
            img.height = size_px
            col_off = int(max((DESIGN_W_PT - IMG_PT) / 2, 1) * EMU_PER_PT)
            img.anchor = OneCellAnchor(
                _from=AnchorMarker(col=1, row=rB - 1, colOff=col_off,
                                   rowOff=int(1.5 * EMU_PER_PT)),
                ext=XDRPositiveSize2D(cx=int(pixels_to_EMU(size_px)),
                                      cy=int(pixels_to_EMU(size_px))))
            ws.add_image(img)

        B_TOP = Border(left=THIN, right=THIN, top=THIN)
        B_MID = Border(left=THIN, right=THIN)
        B_BOT = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        for cc in range(1, NC + 1):
            ws.cell(row=rA, column=cc).border = B_TOP
            ws.cell(row=rB, column=cc).border = B_MID
            ws.cell(row=rC, column=cc).border = B_MID
            ws.cell(row=last, column=cc).border = B_BOT
        page_rows[it['page']] = last
        row = last + 1

    # ---- Add + Total
    add_row = row
    ws.cell(row=add_row, column=22, value='Add').font = F
    ws.cell(row=add_row, column=22).alignment = R
    c = ws.cell(row=add_row, column=23, value=0.0)
    c.font = F
    c.number_format = INR
    c.alignment = R
    ws.row_dimensions[add_row].height = 11
    tot_row = add_row + 1
    ws.cell(row=tot_row, column=2, value='Total').font = FB
    ws.cell(row=tot_row, column=2).alignment = LC

    def ssum(key, src='sub'):
        return round(sum(x[src][key] for x in items
                         if isinstance(x[src].get(key), (int, float))), 3)

    def isum(key):
        return round(sum(x['item'][key] for x in items
                         if isinstance(x['item'].get(key), (int, float))), 3)

    totals = {5: ssum('d_pcs'), 6: ssum('d_wt'), 8: ssum('d_amt'), 10: isum('m_wt'),
              11: isum('m_netwt'), 13: isum('m_amt'), 16: ssum('s_pcs'), 17: ssum('s_wt'),
              19: ssum('s_amt'), 20: ssum('o_amt'), 22: isum('l_amt'), 23: isum('total')}
    for col, v in totals.items():
        cc = ws.cell(row=tot_row, column=col, value=v)
        cc.font = FB
        cc.number_format = FMT_PCS if col in (5, 16) else (FMT_WT if col in (6, 10, 11, 17) else INR)
        cc.alignment = R
    for cc in range(1, NC + 1):
        ws.cell(row=add_row, column=cc).border = BORDER
        ws.cell(row=tot_row, column=cc).border = BORDER
    ws.row_dimensions[tot_row].height = 12

    # ---- summary block
    sr0 = tot_row + 1

    def put(r, c, v, font=F, al=None, fmt=None):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = font
        if al:
            cell.alignment = al
        if fmt:
            cell.number_format = fmt
        return cell

    def box(r1, c1, r2, c2):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).border = BORDER

    ws.merge_cells(start_row=sr0, start_column=1, end_row=sr0, end_column=9)
    put(sr0, 1, 'SUMMARY', FB, C)
    ws.merge_cells(start_row=sr0, start_column=10, end_row=sr0, end_column=13)
    put(sr0, 10, 'Diamond Detail', FB, C)
    ws.merge_cells(start_row=sr0, start_column=14, end_row=sr0, end_column=17)
    put(sr0, 14, 'OTHER DETAILS', FB, C)

    gross_total = 0.0
    for x in items:
        m = re.match(r'([\d.]+)\s*gm', x['gross'] or '')
        if m:
            gross_total += float(m.group(1))
    d_amt = ssum('d_amt')
    s_amt = ssum('s_amt')
    left = [('GOLD IN 24KT', f'{isum("m_wt") * 0.60:.3f} gm'),
            ('GROSS WT', f'{gross_total:.3f} gm'),
            ('*(G+D) WT', f'{isum("m_wt"):.3f} gm'),
            ('NET WT', f'{isum("m_netwt"):.3f} gm'), ('', ''),
            ('DIAMOND WT', f'{int(ssum("d_pcs"))} / {ssum("d_wt"):.3f} cts'),
            ('STONE WT', f'{int(ssum("s_pcs"))} / {ssum("s_wt"):.3f} cts'),
            ('MISC WT', '0 / 0.000 gm')]
    mid = [('GOLD', 0.0), ('DIAMOND', d_amt), ('CST', s_amt), ('MISC', 0.0),
           ('MAKING', isum('l_amt')), ('OTHER', ssum('o_amt')), ('ADD', 0.0),
           ('TOTAL', ssum('total', 'item'))]
    for i, (k, v) in enumerate(left):
        r = sr0 + 1 + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        put(r, 1, k, FB, LC)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        put(r, 3, v, F, Alignment(horizontal='right', vertical='center'))
    for i, (k, v) in enumerate(mid):
        r = sr0 + 1 + i
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        put(r, 6, k, FB, LC)
        ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9)
        put(r, 8, v, F, R, INR)
    r = sr0 + 1
    ws.merge_cells(start_row=r, start_column=10, end_row=r, end_column=11)
    put(r, 10, 'OTHER', FB, LC)
    ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=13)
    put(r, 12, f'{int(ssum("d_pcs"))} / {ssum("d_wt"):.3f} cts', F, LC)
    ws.merge_cells(start_row=r, start_column=14, end_row=r, end_column=15)
    put(r, 14, 'RATE IN 24KT', FB, LC)
    ws.merge_cells(start_row=r, start_column=16, end_row=r, end_column=17)
    put(r, 16, 0.0, F, R, INR)
    box(sr0, 1, sr0 + 8, 9)
    box(sr0, 10, sr0 + 8, 13)
    box(sr0, 14, sr0 + 1, 17)
    for rr in range(sr0, sr0 + 9):
        ws.row_dimensions[rr].height = 12
    cb = sr0 + 10
    ws.merge_cells(start_row=cb, start_column=6, end_row=cb, end_column=9)
    put(cb, 6, 'Created By', FI, C)
    ws.merge_cells(start_row=cb, start_column=14, end_row=cb, end_column=17)
    put(cb, 14, 'Checked By', FI, C)
    for rr in range(sr0 + 2, cb + 1):
        for cc in range(1, 18):
            if rr > sr0 + 1:
                ws.cell(row=rr, column=cc).border = Border()
    for rr in range(sr0, cb + 1):
        for cc in range(18, NC + 1):
            ws.cell(row=rr, column=cc).border = Border()
    ws.row_dimensions[cb].height = 14

    # ---- print setup
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.2
    ws.page_margins.top = ws.page_margins.bottom = 0.25
    ws.print_title_rows = '1:6'
    ws.freeze_panes = 'C7'
    if page_rows:
        mx = max(page_rows)
        for pg, rr in sorted(page_rows.items()):
            if pg != mx:
                ws.row_breaks.append(Break(id=rr))
    wb.save(out_path)
    return {'sheet': ws.title, 'rows': ws.max_row, 'cols': NC}


def verify_estimate(parsed):
    """internal consistency check of every item"""
    items = parsed['items']
    F = lambda v: v if isinstance(v, (int, float)) else 0.0
    bad_total = bad_labour = bad_diamond = 0
    grand = 0.0
    for it in items:
        dia = sum(F(x.get('d_amt')) for x in it['diamonds'])
        sto = sum(F(x.get('s_amt')) for x in it['stones'])
        oth = F(it['sub'].get('o_amt'))
        calc = dia + F(it['item'].get('m_amt')) + sto + oth + F(it['item'].get('l_amt'))
        tot = F(it['item'].get('total')) or F(it['sub'].get('total'))
        grand += tot
        if tot and abs(calc - tot) > 0.06:
            bad_total += 1
        lr, la, nw = F(it['item'].get('l_rate')), F(it['item'].get('l_amt')), F(it['item'].get('m_netwt'))
        if lr and la and nw and abs(lr * nw - la) > 0.4:
            bad_labour += 1
        for x in it['diamonds']:
            w, r, a = F(x.get('d_wt')), F(x.get('d_rate')), F(x.get('d_amt'))
            if w and r and a and abs(w * r - a) > 0.6:
                bad_diamond += 1
    return {'items': len(items), 'grand_total': round(grand, 2),
            'mismatch_item_total': bad_total, 'mismatch_labour': bad_labour,
            'mismatch_diamond': bad_diamond,
            'ok': (bad_total == bad_labour == bad_diamond == 0)}


# =====================================================================
#  GENERIC engine
# =====================================================================


def build_generic_with_images(doc, out_path, img_dir, with_images=True):
    """two pass generic build so images can be anchored to the right row"""
    wb = Workbook()
    ws = wb.active
    ws.title = 'PDF'
    row = 1
    title_rows = None
    breaks = []

    for pno in range(len(doc)):
        page = doc[pno]
        tabs = page.find_tables()
        row_meta = []          # (excel_row, table_row_bbox)
        if not tabs.tables:
            for ln in page.get_text().split('\n'):
                if ln.strip():
                    ws.cell(row=row, column=1, value=clean(ln)).font = F
                    ws.cell(row=row, column=1).alignment = LC
                    ws.row_dimensions[row].height = 12
                    row += 1
            breaks.append(row)
            row += 1
            continue

        for table in tabs.tables:
            edges = edges_of(table)
            if len(edges) < 3:
                continue
            heights = {}
            for tr in table.rows:
                h = tr.bbox[3] - tr.bbox[1]
                heights[row] = max(11.0, h)
                row_meta.append((row, tr.bbox))
                for c in tr.cells:
                    if c is None:
                        continue
                    x0, y0, x1, y1 = c
                    try:
                        c0 = edges.index(round(x0, 1))
                        c1 = edges.index(round(x1, 1)) - 1
                    except ValueError:
                        continue
                    if c1 < c0:
                        c1 = c0
                    txt = fix_wrapped_number(clean(page.get_textbox(pymupdf.Rect(c))))
                    if c1 > c0:
                        ws.merge_cells(start_row=row, start_column=c0 + 1,
                                       end_row=row, end_column=c1 + 1)
                    put_cell(ws, row, c0 + 1, txt, F, LT if not is_num(txt) else R)
                for cc in range(1, len(edges)):
                    ws.cell(row=row, column=cc).border = BORDER
                row += 1
            for i, (tl, tt) in enumerate(
                    zip(range(row - table.row_count, row), [t.bbox for t in table.rows])):
                ws.row_dimensions[tl].height = max(11.0, tt[3] - tt[1])
            for i, w in enumerate([edges[i + 1] - edges[i] for i in range(len(edges) - 1)],
                                  start=1):
                wid = excel_col_width(w)
                cur = ws.column_dimensions[get_column_letter(i)].width or 0
                ws.column_dimensions[get_column_letter(i)].width = max(cur, wid)
            if title_rows is None:
                title_rows = row - table.row_count
            breaks.append(row)
            row += 1

        # ---- images onto the matching rows
        if with_images and page.get_images(full=True):
            for im in collect_page_images(doc, pno, img_dir, ''):
                ix0, iy0, ix1, iy1 = im['bbox']
                # nearest row
                best = None
                for (er, tb) in row_meta:
                    if tb[1] - 3 <= iy0 <= tb[3] + 3:
                        best = (er, tb)
                        break
                if best is None:
                    best = min(row_meta, key=lambda a: abs(a[1][1] - iy0)) if row_meta else None
                if best is None:
                    continue
                er, tb = best
                h_pt = min(iy1 - iy0, 60)
                w_pt = (ix1 - ix0) * (h_pt / max(iy1 - iy0, 0.1))
                # keep inside a sane size
                if w_pt > 120:
                    w_pt = 120
                    h_pt = w_pt * (iy1 - iy0) / max(ix1 - ix0, 0.1)
                need = h_pt + 4
                if (ws.row_dimensions[er].height or 11) < need:
                    ws.row_dimensions[er].height = need
                img = XLImage(im['path'])
                px = lambda pt: pt * 96 / 72
                img.width, img.height = px(w_pt), px(h_pt)
                col0 = 0
                img.anchor = OneCellAnchor(
                    _from=AnchorMarker(col=col0, row=er - 1,
                                       colOff=int(2 * EMU_PER_PT),
                                       rowOff=int(1.5 * EMU_PER_PT)),
                    ext=XDRPositiveSize2D(cx=int(pixels_to_EMU(img.width)),
                                          cy=int(pixels_to_EMU(img.height))))
                ws.add_image(img)

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.2
    ws.page_margins.top = ws.page_margins.bottom = 0.25
    if title_rows:
        pass
    wb.save(out_path)
    return {'rows': ws.max_row, 'cols': ws.max_column}


# =====================================================================
#  TEXT fallback
# =====================================================================


def build_text_xlsx(doc, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'PDF text'
    ws.column_dimensions['A'].width = 120
    r = 1
    for pno in range(len(doc)):
        ws.cell(row=r, column=1, value=f'--- Page {pno + 1} ---').font = FB
        r += 1
        for ln in doc[pno].get_text().split('\n'):
            if ln.strip():
                c = ws.cell(row=r, column=1, value=clean(ln))
                c.font = F
                c.alignment = LC
                r += 1
        r += 1
    wb.save(out_path)
    return {'rows': r}


# =====================================================================
#  images zip
# =====================================================================


def build_images_zip(doc, zip_path, manifest_rows=None):
    seen = {}
    entries = []
    for pno in range(len(doc)):
        page = doc[pno]
        for info in page.get_image_info(xrefs=True):
            xref = info['xref']
            if not xref or xref in seen:
                continue
            data = doc.extract_image(xref)
            if not data or not data.get('image'):
                continue
            h = hashlib.md5(data['image']).hexdigest()[:12]
            name = f'page{pno + 1:02d}_{len(entries) + 1:03d}_{h}.{data["ext"]}'
            seen[xref] = name
            entries.append({'file': name, 'page': pno + 1, 'bytes': len(data['image']),
                            'w': data['width'], 'h': data['height'], 'data': data['image']})
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as z:
        for e in entries:
            z.writestr(e['file'], e['data'])
        if entries:
            import csv
            buf = io.StringIO()
            wr = csv.writer(buf)
            wr.writerow(['file', 'page', 'width_px', 'height_px', 'size_bytes'])
            for e in entries:
                wr.writerow([e['file'], e['page'], e['w'], e['h'], e['bytes']])
            z.writestr('images_list.csv', buf.getvalue())
    return entries


# =====================================================================
#  main entry
# =====================================================================


def process(src_pdf, job_dir, mode='auto', progress=None):
    img_dir = os.path.join(job_dir, 'imgs')
    os.makedirs(img_dir, exist_ok=True)
    doc = pymupdf.open(src_pdf)
    tpl = detect_template(doc)
    used = tpl if mode in ('auto', None) else mode
    if used == 'estimate' and tpl != 'estimate':
        used = tpl
    if used == 'text':
        used = 'generic' if table_check(doc) == 'generic' else 'text'

    info = {'pages': len(doc), 'template': tpl, 'engine': used,
            'file': os.path.basename(src_pdf)}

    with_img = os.path.join(job_dir, 'with_images.xlsx')
    without_img = os.path.join(job_dir, 'without_images.xlsx')

    if used == 'estimate':
        parsed = parse_estimate(doc, img_dir)
        parsed['header'] = parse_header(doc)
        build_estimate_xlsx(parsed, with_img, with_images=True)
        build_estimate_xlsx(parsed, without_img, with_images=False)
        items = parsed['items']
        info['items'] = len(items)
        info['header'] = parsed.get('header') or {}
        info['diamond_lines'] = sum(len(i['diamonds']) for i in items)
        info['stone_lines'] = sum(len(i['stones']) for i in items)
        info['design_codes'] = [i['design'] for i in items]
        info['grand_total'] = sum(i['item']['total'] for i in items
                                  if isinstance(i['item'].get('total'), (int, float)))
        info['total_pcs'] = sum(i['sub']['d_pcs'] for i in items
                                if isinstance(i['sub'].get('d_pcs'), (int, float)))
        info['total_cts'] = sum(i['sub']['d_wt'] for i in items
                                if isinstance(i['sub'].get('d_wt'), (int, float)))
    elif used == 'generic':
        build_generic_with_images(doc, with_img, img_dir, with_images=True)
        build_generic_with_images(doc, without_img, img_dir, with_images=False)
        info['items'] = None
    else:
        build_text_xlsx(doc, with_img)
        shutil.copy(with_img, without_img)
        info['items'] = None

    zip_path = os.path.join(job_dir, 'all_images.zip')
    entries = build_images_zip(doc, zip_path)
    info['images'] = len(entries)
    info['has_images'] = len(entries) > 0

    # preview of the first page
    try:
        pix = doc[0].get_pixmap(dpi=100)
        pix.save(os.path.join(job_dir, 'preview_1.png'))
        info['preview'] = 'preview_1.png'
    except Exception:
        info['preview'] = None

    try:
        info['verify'] = verify_estimate(parsed) if used == 'estimate' else None
    except Exception:
        info['verify'] = None

    info['sizes'] = {
        'with_images': os.path.getsize(with_img),
        'without_images': os.path.getsize(without_img),
        'all_images': os.path.getsize(zip_path),
    }
    doc.close()
    return info
